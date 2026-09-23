import os
import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

PPTX_PATH = 'Adv GLE_IBM BOB_Project Submission Template for EduentFoundation_IBMUniversityEngagement11 (1).pptx'

def customize_presentation():
    print(f"Loading presentation: {PPTX_PATH}")
    prs = pptx.Presentation(PPTX_PATH)

    # ----------------------------------------------------
    # Helper: replace image blob of shape
    # ----------------------------------------------------
    def replace_image(shape, image_path):
        if not os.path.exists(image_path):
            print(f"Warning: image {image_path} not found!")
            return
        img_part = shape.part.related_part(shape._element.blip_rId)
        with open(image_path, 'rb') as f:
            img_part._blob = f.read()
        print(f"Replaced image in {shape.name} with {image_path}")

    # ====================================================
    # SLIDE 1: Title Slide
    # ====================================================
    s1 = prs.slides[0]
    for shape in s1.shapes:
        if shape.name == 'PlaceHolder 1':
            shape.text_frame.text = "IBM University Engagement Project Submission"
            for p in shape.text_frame.paragraphs:
                p.font.name = 'Arial'
                p.font.bold = True
                p.font.size = Pt(28)
        elif shape.name == 'PlaceHolder 2':
            shape.text_frame.text = (
                "“Problem Statement No. 7: Research Agent”\n"
                "Autonomous Multi-Agent Academic Discovery & Literature Synthesis Studio\n"
                "Powered by IBM Granite 3.0, IBM watsonx.ai, and IBM Bob\n\n"
                "Author: Vaibhav Chavanpatil"
            )
            for p in shape.text_frame.paragraphs:
                p.font.name = 'Arial'
                p.font.size = Pt(16)
    print("Slide 1 updated.")

    # ====================================================
    # SLIDE 2: Project Details & Student Profile
    # ====================================================
    s2 = prs.slides[1]
    for shape in s2.shapes:
        if shape.name == 'TextBox 6':
            shape.text_frame.text = "Project Title – Problem Statement No. 7: Research Agent (Langflow Multi-Agent Studio)"
            for p in shape.text_frame.paragraphs:
                p.font.name = 'Arial'
                p.font.bold = True
                p.font.size = Pt(20)
        elif shape.name == 'TextBox 3':
            shape.text_frame.text = "Domain of Project – Higher Education & Academic Research / Agentic AI"
            for p in shape.text_frame.paragraphs:
                p.font.name = 'Arial'
                p.font.bold = True
                p.font.size = Pt(20)
        elif shape.name == 'TextBox 2':
            # Clear the recommended font instructions box
            shape.text_frame.text = ""
        elif shape.has_table:
            table = shape.table
            # Row 1 values: Student Name, Photo (shape 3), Email, WhatsApp
            cell_name = table.cell(1, 0)
            cell_name.text = "Vaibhav Chavanpatil\n(3rd Year CSE Student)"
            for p in cell_name.text_frame.paragraphs:
                p.font.name = 'Arial'
                p.font.size = Pt(13)
                p.font.bold = True

            cell_photo = table.cell(1, 1)
            cell_photo.text = "" # Keep cell text empty so picture shows cleanly

            cell_email = table.cell(1, 2)
            cell_email.text = "91275740+VAIBHAV7848@users.noreply.github.com\nGitHub: github.com/VAIBHAV7848"
            for p in cell_email.text_frame.paragraphs:
                p.font.name = 'Arial'
                p.font.size = Pt(11)

            cell_phone = table.cell(1, 3)
            cell_phone.text = "+91 91275 74800\n[Available on WhatsApp]"
            for p in cell_phone.text_frame.paragraphs:
                p.font.name = 'Arial'
                p.font.size = Pt(12)
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.name == 'Google Shape;98;p2':
            # Replace passport photo with user avatar
            replace_image(shape, 'avatar.png')
    print("Slide 2 updated.")

    # ====================================================
    # SLIDE 3: Problem Statement
    # ====================================================
    s3 = prs.slides[2]
    for shape in s3.shapes:
        if shape.name == 'TextBox 1':
            shape.text_frame.text = (
                "Problem Statement : Research Agent (Problem Statement No. 7)\n\n"
                "The Challenge\n"
                "Academic researchers, scientists, and university students grapple with literature overload across tens of thousands of preprint publications annually. Critical scientific insights, empirical benchmark comparisons, and citation gaps remain buried in siloed PDFs. Manual synthesis is labor-intensive, slow, ungrounded, and prone to severe hallucinations in generic LLM chatbots.\n\n"
                "The Objective\n"
                "Build an autonomous, agentic Research Agent platform powered by IBM Bob, IBM watsonx.ai, and IBM Granite 3.0 that delivers grounded, multi-agent academic discovery:\n"
                "• 2D Interactive Physics Citation Topology – Real-time force-directed clustering of papers and citation lineage at 60 FPS.\n"
                "• AST Academic Document Ingestion – Automated parser extracting metadata, abstracts, and reference links into the active corpus.\n"
                "• Emerging Research Velocity & Trajectories – Predictive momentum forecasting (2021–2026) across domain datasets.\n"
                "• Citation Gap & White-Space Matrix – Algorithmic detection of unexplored intersections and predictive hypothesis generation.\n"
                "• 4-Agent Langflow Pipeline – Collaborative orchestration of AST Parser, Topology Builder, Gap Detector, and Granite RAG Reasoner.\n"
                "• Automated Literature Review & Export – Instant generation of peer-review-ready synthesis with BibTeX (.bib) and Markdown (.md) export."
            )
            for i, p in enumerate(shape.text_frame.paragraphs):
                p.font.name = 'Arial'
                p.font.size = Pt(10)
        elif shape.name == 'TextBox 2':
            shape.text_frame.text = "Problem Statement -"
    print("Slide 3 updated.")

    # ====================================================
    # SLIDE 4: Proposed Solution
    # ====================================================
    s4 = prs.slides[3]
    for shape in s4.shapes:
        if shape.name == 'TextBox 1':
            shape.text_frame.text = (
                "Proposed Solution: AI-Powered Research Agent Studio\n\n"
                "The proposed solution is an autonomous multi-agent academic discovery and literature review studio built using IBM Bob, IBM watsonx.ai, and IBM Granite 3.0 models. The system addresses challenges in literature comprehension, cross-paper comparison, and white-space discovery through a modular, agentic architecture.\n\n"
                "At its core, the solution utilizes a Graph-Grounded Retrieval-Augmented Generation (RAG) pipeline to fetch accurate context from peer-reviewed academic corpora. Citations and claims are embedded into a dynamic topological graph, enabling the system to evaluate empirical metrics and comparative trade-offs using Granite models.\n\n"
                "The system consists of four specialized collaborative agents:\n"
                "• AST Ingestion & Parser Agent: Extracts academic metadata, abstracts, and reference citations from uploaded PDFs and preprints.\n"
                "• Topology & Citation Graph Agent: Calculates connection strengths, thematic clusters, and force-directed 2D physics layouts.\n"
                "• Gap Detector & Trend Velocity Agent: Identifies unexplored literature intersections, computes opportunity scores (0-100), and models publication trajectories.\n"
                "• Granite RAG Synthesis Agent: Reasons across citations, answers complex cross-paper questions, and writes comprehensive literature reviews with verified attribution.\n\n"
                "The solution features a high-performance, human-centered studio interface (Claude Terracotta & Warm White design system) providing researchers with real-time graph navigation, trajectory charts, slide-over paper drawers, and BibTeX/Markdown export.\n"
                "IBM Bob orchestrates multi-agent workflows, while watsonx.ai ensures enterprise-grade model deployment, governance, and low-latency reasoning with Groq LPU acceleration."
            )
            for p in shape.text_frame.paragraphs:
                p.font.name = 'Arial'
                p.font.size = Pt(10)
        elif shape.name == 'TextBox 2':
            shape.text_frame.text = "Proposed Solution -"
    print("Slide 4 updated.")

    # ====================================================
    # SLIDE 5: Technology Used
    # ====================================================
    s5 = prs.slides[4]
    for shape in s5.shapes:
        if shape.name == 'TextBox 1':
            shape.text_frame.text = "Technology Used"
        elif shape.name == 'TextBox 2':
            shape.text_frame.text = (
                "IBM Bob Platform – Visual workflow design, multi-agent orchestration, and full-stack development environment for AI agent systems.\n\n"
                "IBM Granite 3.0 Models (ibm/granite-3-8b-instruct) – Enterprise foundation models for natural language understanding, reasoning, tool invocation, and academic literature synthesis.\n\n"
                "IBM Cloud & watsonx.ai – Provides the scalable enterprise infrastructure to build, deploy, and govern the Research Agent with vector embeddings.\n\n"
                "Groq LPU Inference Engine – Ultra-fast live LPU execution (~500 T/s) with resilient offline fallback for zero-downtime academic demonstrations.\n\n"
                "Graph-Grounded RAG – Retrieves verified paper context and citation topology before generating answers, eliminating hallucinations.\n\n"
                "Next.js 15, React 19 & Tailwind CSS – High-performance studio interface with custom 2D Physics Canvas and Claude Terracotta design."
            )
            for p in shape.text_frame.paragraphs:
                p.font.name = 'Arial'
                p.font.size = Pt(12)
    print("Slide 5 updated.")

    # ====================================================
    # SLIDE 6: Required Files created using IBM bob
    # ====================================================
    s6 = prs.slides[5]
    for shape in s6.shapes:
        if shape.name == 'Google Shape;135;p5':
            shape.text_frame.text = "Required Files created using IBM bob -"
        elif shape.name == 'TextBox 1':
            shape.text_frame.text = (
                "app/page.tsx\n\n"
                "CitationCanvas.tsx\n\n"
                "TrendVelocityView.tsx\n\n"
                "CitationGapMatrix.tsx\n\n"
                "GraniteChatConsole.tsx\n\n"
                "LiteratureReviewView.tsx\n\n"
                "app/api/chat/route.ts\n\n"
                "package.json / .env.local"
            )
            for p in shape.text_frame.paragraphs:
                p.font.name = 'Arial'
                p.font.size = Pt(11)
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.name == 'Picture 4':
            replace_image(shape, 'screenshots/ibm_bob_ide_screenshot.png')
    print("Slide 6 updated.")

    # ====================================================
    # SLIDE 7: IBM Bob Screenshots (Initial Studio)
    # ====================================================
    s7 = prs.slides[6]
    for shape in s7.shapes:
        if shape.name == 'Google Shape;135;p5':
            shape.text_frame.text = "IBM Bob Screenshots: 2D Physics Citation Graph & Corpus Library"
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.name == 'Picture 4':
            replace_image(shape, 'screenshots/01_initial_studio.png')
    print("Slide 7 updated.")

    # ====================================================
    # SLIDE 8: IBM Bob Screenshots (Chat & Drawer)
    # ====================================================
    s8 = prs.slides[7]
    for shape in s8.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.name == 'Picture 4':
            replace_image(shape, 'screenshots/06_chat_synthesis.png')
    print("Slide 8 updated.")

    # ====================================================
    # SLIDE 9: Architecture Blueprint
    # ====================================================
    s9 = prs.slides[8]
    for shape in s9.shapes:
        if shape.name == 'Google Shape;135;p5':
            shape.text_frame.text = "Architecture blueprint for IBM bob: 4-Agent Langflow & Granite RAG"
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.name == 'Picture 2':
            replace_image(shape, 'screenshots/architecture_blueprint.png')
    print("Slide 9 updated.")

    # ====================================================
    # SLIDE 10: Role of Agentic AI in the solution
    # ====================================================
    s10 = prs.slides[9]
    for shape in s10.shapes:
        if shape.name == 'TextBox 1':
            shape.text_frame.text = "Role of Agentic AI in the solution"
        elif shape.name == 'TextBox 2':
            shape.text_frame.text = (
                "Role of Agentic AI in Research Agent Studio\n\n"
                "Agentic AI elevates the Research Agent from a passive question-answering tool into an autonomous, proactive research co-pilot. Powered by IBM watsonx.ai and IBM Granite 3.0, it coordinates four specialized agents acting in concert:\n\n"
                "Each agent performs a specific role—such as parsing academic preprints, constructing citation topologies, forecasting publication velocity, and synthesizing cross-paper literature reviews—while collaborating autonomously.\n\n"
                "Agentic AI brings deep context awareness, understanding user domain selections (e.g. Agentic AI, Quantum Computing, or Biomedical AI) to retrieve grounded evidence. It continuously adapts when new PDFs are ingested into the active graph.\n\n"
                "Additionally, agents communicate with each other through structured contracts, enabling automated hypothesis formulation and peer-review synthesis.\n\n"
                "Overall, Agentic AI makes the system proactive, goal-driven, and verifiable, acting as a tireless digital researcher that compresses weeks of manual literature review into seconds."
            )
            for p in shape.text_frame.paragraphs:
                p.font.name = 'Arial'
                p.font.size = Pt(11)
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            replace_image(shape, 'screenshots/slide10_agentic_ai.png')
    print("Slide 10 updated.")

    # ====================================================
    # SLIDE 11: Project Output Screenshot - 2D Graph
    # ====================================================
    s11 = prs.slides[10]
    for shape in s11.shapes:
        if shape.name == 'Google Shape;278;p18':
            shape.text_frame.text = "Project Output Screenshot – 2D Physics Citation Topology"
        elif shape.name == 'TextBox 1':
            shape.text_frame.text = "Interactive 2D force-directed physics graph rendering cross-paper citations, thematic clusters, node halos, and metadata cards at 60 FPS."
            for p in shape.text_frame.paragraphs:
                p.font.name = 'Arial'
                p.font.size = Pt(11)
        elif shape.name in ['Google Shape;280;p18', 'Google Shape;281;p18', 'Google Shape;282;p18']:
            shape.text_frame.text = ""
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.name == 'Picture 3':
            replace_image(shape, 'screenshots/01_initial_studio.png')
    print("Slide 11 updated.")

    # ====================================================
    # SLIDE 12: Project Output Screenshot - Trends
    # ====================================================
    s12 = prs.slides[11]
    for shape in s12.shapes:
        if shape.name == 'Google Shape;278;p18':
            shape.text_frame.text = "Project Output Screenshot – Emerging Trajectories & Velocity"
        elif shape.name == 'TextBox 1':
            shape.text_frame.text = "Predictive modeling of 5-year academic publication momentum (2021-2026) alongside milestone catalyst paper identification."
            for p in shape.text_frame.paragraphs:
                p.font.name = 'Arial'
                p.font.size = Pt(11)
        elif shape.name in ['Google Shape;280;p18', 'Google Shape;281;p18', 'Google Shape;282;p18']:
            shape.text_frame.text = ""
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.name == 'Picture 4':
            replace_image(shape, 'screenshots/03_emerging_trends.png')
    print("Slide 12 updated.")

    # ====================================================
    # SLIDE 13: Project Output Screenshot - Literature Review
    # ====================================================
    s13 = prs.slides[12]
    for shape in s13.shapes:
        if shape.name == 'Google Shape;278;p18':
            shape.text_frame.text = "Project Output Screenshot – Automated Literature Review Synthesis"
        elif shape.name == 'TextBox 1':
            shape.text_frame.text = "Granite 3.0 multi-agent synthesis generating structured executive summaries, benchmark comparison tables, and one-click BibTeX/Markdown export."
            for p in shape.text_frame.paragraphs:
                p.font.name = 'Arial'
                p.font.size = Pt(11)
        elif shape.name in ['Google Shape;280;p18', 'Google Shape;281;p18', 'Google Shape;282;p18']:
            shape.text_frame.text = ""
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.name == 'Picture 3':
            replace_image(shape, 'screenshots/07_literature_review_generated.png')
    print("Slide 13 updated.")

    # ====================================================
    # SLIDE 14: Novelty and Uniqueness
    # ====================================================
    s14 = prs.slides[13]
    for shape in s14.shapes:
        if shape.name == 'TextBox 1':
            shape.text_frame.text = "Novelty and Uniqueness"
        elif shape.name == 'TextBox 2':
            shape.text_frame.text = (
                "The Langflow Research Agent Studio sets a new benchmark by fusing Agentic AI, real-time physics graphs, and Granite RAG into a unified scientific workspace.\n\n"
                "• Multi-Agent Intelligence – Four specialized collaborative agents (AST Ingestion, Topology Builder, Gap Detector, Granite Reasoner) work in concert.\n\n"
                "• 2D Physics Citation Topology – Interactive force-directed canvas visualizing paper clusters, connection strengths, and reference networks at 60 FPS.\n\n"
                "• Algorithmic Citation Gap Detection – Automatically identifies literature white spaces, calculates Opportunity Scores (0-100), and drafts testable hypotheses.\n\n"
                "• Publication Velocity Forecasting – Empirical modeling of 5-year academic publication momentum (2021-2026) across domain corpora.\n\n"
                "• Dual-Engine Live AI (Groq LPU + IBM Granite) – Sub-second inference (~500 T/s) combined with enterprise reasoning and zero-config offline fallback.\n\n"
                "• Iconic Claude Terracotta Aesthetics – Warm paper cream background (#FAF9F5), terracotta orange accents, and slide-over paper detail drawer.\n\n"
                "• Enterprise-Ready Export – Formats citations into publication-standard BibTeX (.bib) and comprehensive Markdown (.md) with verified attribution."
            )
            for p in shape.text_frame.paragraphs:
                p.font.name = 'Arial'
                p.font.size = Pt(11)
    print("Slide 14 updated.")

    # ====================================================
    # SLIDE 15: GitHub Link
    # ====================================================
    s15 = prs.slides[14]
    for shape in s15.shapes:
        if shape.name == 'Google Shape;278;p18':
            shape.text_frame.text = "Git Hub Link"
        elif shape.name == 'Google Shape;279;p18':
            shape.text_frame.text = (
                "GitHub Repository & Submission Deliverables:\n\n"
                "• Public GitHub Repository URL:\n"
                "  https://github.com/VAIBHAV7848/IBM_WORKSHOP\n\n"
                "• Author: Vaibhav Chavanpatil (GitHub: @VAIBHAV7848)\n"
                "• Project: Problem Statement No. 7 – Research Agent Studio\n\n"
                "Uploaded Deliverables in Repository:\n"
                "1) Complete Next.js & IBM Granite Multi-Agent Source Code (components, physics engine, API routes)\n"
                "2) 7 Problem statement on Langflow 2026-2027.pdf (Official Challenge Documentation)\n"
                "3) Adv GLE_IBM BOB_Project Submission Presentation.pptx (Complete Presentation Deck)\n"
                "4) Live Groq LPU Integration, Verification Suite, and High-Res Demo Screenshots\n"
                "5) Comprehensive Setup & Quick Start Guide with Zero-Config Offline Fallback"
            )
            for p in shape.text_frame.paragraphs:
                p.font.name = 'Arial'
                p.font.size = Pt(12)
        elif shape.name in ['Google Shape;280;p18', 'Google Shape;281;p18', 'Google Shape;282;p18']:
            shape.text_frame.text = ""
    print("Slide 15 updated.")

    # ====================================================
    # SLIDE 16: Future Scope
    # ====================================================
    s16 = prs.slides[15]
    for shape in s16.shapes:
        if shape.name == 'Google Shape;278;p18':
            shape.text_frame.text = "Future Scope"
        elif shape.name == 'Google Shape;279;p18':
            shape.text_frame.text = (
                "1. Live Academic API Synchronization (arXiv, Semantic Scholar, CrossRef) –\n"
                "The system can be integrated with arXiv and Semantic Scholar APIs to continuously ingest newly published preprints in real time, automatically updating citation topology graphs and velocity trends as new research drops.\n\n"
                "2. Graph Neural Network (GNN) Link Prediction –\n"
                "Deploying inductive Graph Convolutional Networks (GCN) on top of the citation topology to predict emerging cross-disciplinary breakthroughs and co-authorship opportunities before they appear in mainstream venues.\n\n"
                "3. Multi-Modal LaTeX Formula & Benchmark Extraction –\n"
                "Enhancing the AST parser with Vision-Language Models to extract mathematical equations, loss functions, and benchmark tables directly into computational notebooks.\n\n"
                "4. Federated Multi-Institution Knowledge Mesh –\n"
                "Enabling universities and research labs to host decentralized citation nodes while preserving proprietary preprint confidentiality through IBM Granite Guardian safety guardrails."
            )
            for p in shape.text_frame.paragraphs:
                p.font.name = 'Arial'
                p.font.size = Pt(11)
        elif shape.name in ['Google Shape;280;p18', 'Google Shape;281;p18', 'Google Shape;282;p18']:
            shape.text_frame.text = ""
    print("Slide 16 updated.")

    # ====================================================
    # SLIDE 17 & 18: Certificates Earned
    # ====================================================
    s17 = prs.slides[16]
    for shape in s17.shapes:
        if shape.name == 'TextBox 2':
            shape.text_frame.text = "IBM SkillsBuild & Credly Verified Credentials\nStudent: Vaibhav Chavanpatil"
            for p in shape.text_frame.paragraphs:
                p.font.name = 'Arial'
                p.font.size = Pt(14)
                p.font.bold = True
    print("Slide 17 updated.")

    s18 = prs.slides[17]
    for shape in s18.shapes:
        if shape.name == 'TextBox 5':
            shape.text_frame.text = "IBM BOB Workshop Participation Certificate\nStudent: Vaibhav Chavanpatil"
            for p in shape.text_frame.paragraphs:
                p.font.name = 'Arial'
                p.font.size = Pt(14)
                p.font.bold = True
    print("Slide 18 updated.")

    # Save modified presentation
    prs.save(PPTX_PATH)
    print("✅ Successfully updated presentation without altering template structure!")

if __name__ == '__main__':
    customize_presentation()
